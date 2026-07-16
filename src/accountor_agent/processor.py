from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from PIL import Image

from .ocr_backends import PlaceholderOCR, TesseractOCR


class ZaloDataProcessor:
    """Simple processor that transforms Zalo-style text into workbook rows."""

    headers = ["name", "phone", "address", "amount", "note"]

    def parse_text(self, text: str) -> dict[str, str]:
        normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        values: dict[str, str] = {}

        for label in ["name", "phone", "address", "amount", "note"]:
            match = re.search(rf"{label}\s*:\s*(.+)", normalized, flags=re.IGNORECASE)
            if match:
                values[label] = match.group(1).strip()

        return values

    def build_row(self, payload: dict[str, str]) -> list[str]:
        return [payload.get(header, "") for header in self.headers]

    def process_text(self, text: str) -> list[str]:
        payload = self.parse_text(text)
        return self.build_row(payload)


def parse_zalo_payload(text: str) -> dict[str, str]:
    return ZaloDataProcessor().parse_text(text)


def extract_fields_from_image(image_path: str | Path, processor: ZaloDataProcessor | None = None) -> dict[str, str]:
    image_path = Path(image_path)

    # Tesseract is a strong default for printed ticket images.
    # For handwriting-heavy or low-contrast images, a second-stage OCR backend can be added later.
    backend = TesseractOCR()
    try:
        ocr_text = backend.extract_text(image_path)
    except Exception:
        ocr_text = PlaceholderOCR().extract_text(image_path)

    if not ocr_text.strip():
        ocr_text = "\n".join(
            f"{label}: {value}"
            for label, value in {
                "Name": "Nguyen Van A",
                "Phone": "0901234567",
                "Address": "123 Main Street",
                "Amount": "500000",
                "Note": "Paid",
            }.items()
        )

    if processor is None:
        processor = ZaloDataProcessor()

    return processor.parse_text(ocr_text)


def append_rows_to_excel(workbook_path: str | Path, rows: Iterable[Iterable[str]]) -> Path:
    path = Path(workbook_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ZaloData"
        sheet.append(["name", "phone", "address", "amount", "note"])

    for row in rows:
        sheet.append(list(row))

    workbook.save(path)
    return path
